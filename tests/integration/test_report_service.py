from pathlib import Path

import pytest
from docx import Document

from aar.ingestion.commander_docx import generate_commander_report_template
from aar.models import CommanderReport, ParseStatus, SourceDocument
from aar.services.report_service import DuplicateReportError, ingest_commander_report


def _make_filled_docx_bytes(tmp_path: Path, callsign: str = "P12") -> bytes:
    path = tmp_path / f"{callsign}.docx"
    generate_commander_report_template(path)
    doc = Document(str(path))
    doc.tables[0].rows[0].cells[1].text = callsign
    doc.tables[0].rows[3].cells[1].text = "Janssens"
    doc.tables[1].rows[0].cells[1].text = "Ja"
    doc.save(str(path))
    return path.read_bytes()


def test_ingest_commander_report_success(tmp_path, session, sample_incident, monkeypatch):
    monkeypatch.setattr(
        "aar.services.report_service.get_settings",
        lambda: type("S", (), {"project_root": tmp_path})(),
    )
    file_bytes = _make_filled_docx_bytes(tmp_path)
    doc = ingest_commander_report(session, sample_incident, file_bytes, "p12.docx")

    assert doc.parse_status == ParseStatus.PARSED
    report = session.query(CommanderReport).filter_by(source_document_id=doc.id).one()
    assert report.callsign == "P12"
    assert report.surname == "Janssens"


def test_ingest_duplicate_report_raises(tmp_path, session, sample_incident, monkeypatch):
    monkeypatch.setattr(
        "aar.services.report_service.get_settings",
        lambda: type("S", (), {"project_root": tmp_path})(),
    )
    file_bytes = _make_filled_docx_bytes(tmp_path)
    ingest_commander_report(session, sample_incident, file_bytes, "p12.docx")

    with pytest.raises(DuplicateReportError):
        ingest_commander_report(session, sample_incident, file_bytes, "p12_copy.docx")

    assert session.query(SourceDocument).count() == 1


def test_ingest_corrupt_file_records_error_status(session, sample_incident, tmp_path, monkeypatch):
    monkeypatch.setattr(
        "aar.services.report_service.get_settings",
        lambda: type("S", (), {"project_root": tmp_path})(),
    )
    doc = ingest_commander_report(session, sample_incident, b"not a docx", "broken.docx")
    assert doc.parse_status == ParseStatus.ERROR
    assert doc.parse_error
    assert session.query(CommanderReport).filter_by(source_document_id=doc.id).count() == 0
